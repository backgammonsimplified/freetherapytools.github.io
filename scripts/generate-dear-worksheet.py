"""Generate original Free Therapy Tools practice downloads (python-docx/reportlab).

Run from the repository root. Source handout transcriptions are deliberately not inputs.
"""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import simpleSplit

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "site/resources/interpersonal-effectiveness"
TITLE = "DEAR MAN Script Worksheet"
PAGES = [
    ("Prepare what I want to say", [
        ("Situation", "What is going on?", 4),
        ("Objective", "What result do I want?", 4),
        ("Describe", "What are the observable facts?", 4),
        ("Express", "What do I feel or think?", 4),
    ]),
    ("Make the request and plan my approach", [
        ("Assert", "What am I clearly asking for?", 4),
        ("Reinforce", "Why would cooperation help?", 4),
        ("Mindful", "What phrase will I return to if the conversation gets pulled off track?", 4),
        ("Appear Confident", "How do I want to present myself?", 4),
        ("Negotiate", "What workable alternatives could I offer?", 4),
    ]),
    ("Rehearse and reflect", [
        ("Final script", "How could the pieces sound together when I say them aloud?", 12),
        ("Reflection after using the script", "What worked?", 4),
        ("Next time", "What would I change next time?", 4),
    ]),
]
INTRO = "Use this worksheet when a concrete result is your main priority. Write a clear, respectful request in your own words. Keep the facts and benefits truthful, and rehearse a script you could say aloud."


def make_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.75)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05
    for name, size in [("Title", 23), ("Heading 1", 15), ("Heading 2", 11)]:
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Arial", Pt(size), RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(3)
    for number, (heading, fields) in enumerate(PAGES, 1):
        if number > 1: doc.add_page_break()
        if number == 1:
            doc.add_paragraph(TITLE, "Title")
            doc.add_paragraph("Free Therapy Tools | Original practice worksheet")
            doc.add_paragraph(INTRO)
        doc.add_paragraph(heading, "Heading 1")
        for label, prompt, count in fields:
            doc.add_paragraph(label, "Heading 2")
            p = doc.add_paragraph(prompt)
            p.paragraph_format.keep_with_next = True
            for _ in range(count):
                line = doc.add_paragraph(" ")
                line.paragraph_format.space_after = Pt(0)
                line.paragraph_format.line_spacing = Pt(19)
                borders = OxmlElement("w:pBdr")
                border = OxmlElement("w:bottom")
                for key, value in [("val", "single"), ("sz", "3"), ("color", "D9D9D9")]:
                    border.set(qn("w:" + key), value)
                borders.append(border)
                line._p.get_or_add_pPr().append(borders)
    footer = sec.footer.paragraphs[0]
    footer.text = "Free Therapy Tools   |   DEAR MAN Script Worksheet   |   "
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.style.font.size = Pt(8)
    props = doc.core_properties
    props.title, props.author = TITLE, "Free Therapy Tools"
    props.created = props.modified = datetime(2026, 9, 6, tzinfo=timezone.utc)
    doc.save(path)


def make_pdf(path):
    pdf = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    pdf.setTitle(TITLE); pdf.setAuthor("Free Therapy Tools")
    for number, (heading, fields) in enumerate(PAGES, 1):
        y = 744
        if number == 1:
            pdf.setFont("Helvetica-Bold", 23); pdf.drawString(54, y, TITLE); y -= 23
            pdf.setFont("Helvetica", 10); pdf.drawString(54, y, "Free Therapy Tools | Original practice worksheet"); y -= 22
            for line in simpleSplit(INTRO, "Helvetica", 10, 504):
                pdf.drawString(54, y, line); y -= 13
            y -= 9
        pdf.setFont("Helvetica-Bold", 15); pdf.drawString(54, y, heading); y -= 29
        for label, prompt, count in fields:
            pdf.setFont("Helvetica-Bold", 11); pdf.drawString(54, y, label); y -= 15
            pdf.setFont("Helvetica", 10)
            for line in simpleSplit(prompt, "Helvetica", 10, 504):
                pdf.drawString(54, y, line); y -= 13
            pdf.setStrokeColorRGB(.82, .82, .82); pdf.setLineWidth(.4)
            for _ in range(count):
                y -= 19; pdf.line(54, y, 558, y)
            y -= 21
        assert y >= 45, (number, y)
        pdf.setFont("Helvetica", 8)
        pdf.drawString(54, 27, "Free Therapy Tools | DEAR MAN Script Worksheet")
        pdf.drawRightString(558, 27, f"{number} / 3")
        pdf.showPage()
    pdf.save()


if __name__ == "__main__":
    OUTPUT.mkdir(parents=True, exist_ok=True)
    make_docx(OUTPUT / "dear-man-script-worksheet.docx")
    make_pdf(OUTPUT / "dear-man-script-worksheet.pdf")
    print("Generated original DEAR MAN worksheet DOCX and PDF (3 planned pages each).")
